"""Supporting Excel workbook: index, definitions, source register, observations, metrics, chart data, revisions, quality."""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import config
from ..storage.db import Database

HEADER_FILL = PatternFill("solid", fgColor="6F00B6")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, name: str, header: list[str], rows: list[list[Any]], widths: list[int] | None = None):
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def build_workbook(fp: dict[str, Any], nar: dict[str, Any], db: Database, out_path: Path, metrics_table: pd.DataFrame | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    ws["A1"] = config.term("report_title", fp.get("lang", "en"))
    ws["A1"].font = Font(bold=True, size=14, color="6F00B6")
    ws["A2"] = f"Edition {fp['edition'].get('edition_month')} · as-of {fp['as_of']} (Asia/Baku) · generated {fp['generated_at']} · fact pack {fp.get('fact_pack_hash')}"
    ws["A3"] = f"Information-set mode: {fp.get('information_set_mode')} · narrative mode: {nar.get('mode')}"
    ws["A5"] = "Sheets"
    ws["A5"].font = Font(bold=True)
    sheets = [("Scorecard", "M03 scorecard rows with periods and changes"), ("Definitions", "Metric dictionary (formulas, inputs, units, basis)"),
              ("SourceRegister", "Datasets, document URLs, publication dates, status"), ("Documents", "Downloaded documents with hashes and retrieval times"),
              ("Observations", "Current-vintage observations from history start (with provenance)"), ("Metrics", "Computed metric series"),
              ("ChartData", "Exact data behind every chart, by slide"), ("KPIs", "All headline values shown on slides"), ("Revisions", "Vintages with revised observations"),
              ("Quality", "Data-quality checks"), ("Availability", "Slide-input availability matrix"), ("Narrative", "Findings and slide texts with classification and validation")]
    for i, (n, desc) in enumerate(sheets, start=6):
        ws[f"A{i}"] = n
        ws[f"B{i}"] = desc
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    # Scorecard
    rows = []
    for r in fp["slides"]["M03"]["rows"]:
        rows.append([r["key"], r["row_label"], r["id"], json.dumps(r.get("dims") or {}), r["latest"]["value"] if r.get("latest") else None, r["latest"]["period"] if r.get("latest") else None,
                     r["prior"]["value"] if r.get("prior") else None, r["prior"]["period"] if r.get("prior") else None, r.get("change"), r.get("unit"), r.get("compare"), r.get("period_type")])
    _sheet(wb, "Scorecard", ["key", "label", "metric_id", "dims", "latest", "latest_period", "prior", "prior_period", "change", "unit", "comparison", "period_type"], rows, [16, 40, 30, 20, 12, 12, 12, 12, 12, 8, 12, 24])

    # Definitions
    rows = [[d["id"], d.get("label"), d["formula"], json.dumps(d.get("inputs"), ensure_ascii=False), d.get("unit"), d.get("basis")] for d in fp["definitions"]]
    _sheet(wb, "Definitions", ["metric_id", "label", "formula", "inputs", "unit", "basis"], rows, [38, 50, 18, 60, 10, 60])

    # Source register
    reg = fp["slides"]["A02"]["register"]
    rows = [[r["source_id"], r["institution"], r["dataset_id"], r.get("title_en"), r.get("title_original"), r.get("role"), r.get("discovery_url"), r.get("document_url"), r.get("format"),
             r.get("published_at"), r.get("published_at_basis"), r.get("frequency"), r.get("expected_lag_days"), r.get("parser"), r.get("status"), r.get("latest_period_end"), r.get("n_documents"),
             r.get("history_in_file"), r.get("known_limitations")] for r in reg]
    _sheet(wb, "SourceRegister", ["source_id", "institution", "dataset_id", "title_en", "title_original", "role", "discovery_url", "document_url", "format", "published_at", "published_at_basis",
                                  "frequency", "expected_lag_days", "parser", "status", "latest_period_end", "n_documents", "history_in_file", "known_limitations"], rows, [14, 10, 26, 45, 45, 10, 40, 50, 6, 12, 14, 8, 8, 22, 18, 14, 8, 8, 30])

    # Documents
    rows = [[s["doc_id"], s["source_id"], s["dataset_id"], s.get("title_original"), s.get("document_url"), s.get("published_at"), s.get("published_at_basis"), s.get("retrieved_at"), s.get("sha256"), s.get("size_bytes"), s.get("status")]
            for s in fp["sources"]]
    _sheet(wb, "Documents", ["doc_id", "source_id", "dataset_id", "title_original", "document_url", "published_at", "published_at_basis", "retrieved_at", "sha256", "size_bytes", "status"], rows, [34, 14, 26, 50, 55, 12, 14, 20, 66, 10, 10])

    # Observations (history window, current vintage)
    start = config.settings().get("history_start", "2020-01") + "-01"
    obs = [dict(r) for r in db.current_observations()]
    cols = ["series_id", "dims", "period_start", "period_end", "freq", "period_type", "value", "value_raw", "missing_reason", "unit", "scale_note", "population", "basis", "source_id", "dataset_id", "doc_id", "sheet", "cell_ref",
            "label_original", "extraction_method", "vintage_id", "first_observed_at", "published_at", "flags"]
    rows = [[o.get(c) for c in cols] for o in obs if (o.get("period_end") or "") >= start]
    _sheet(wb, "Observations", cols, rows, [34, 20, 11, 11, 5, 22, 14, 14, 12, 10, 8, 18, 30, 12, 26, 34, 10, 14, 30, 10, 40, 20, 12, 20])

    # Metrics
    if metrics_table is not None and not metrics_table.empty:
        mt = metrics_table[metrics_table["period_end"].astype(str) >= start]
        rows = mt[["metric_id", "dims", "period_end", "value", "unit", "formula", "period_type"]].values.tolist()
        _sheet(wb, "Metrics", ["metric_id", "dims", "period_end", "value", "unit", "formula", "period_type"], [[str(v) if isinstance(v, dt.date) else v for v in r] for r in rows], [40, 22, 12, 14, 8, 18, 26])

    # Chart data
    rows = []
    for sid, sl in fp["slides"].items():
        for k, v in (sl or {}).items():
            if isinstance(v, list):
                for item in v:
                    if isinstance(item, dict) and "points" in item:
                        for p in item["points"]:
                            rows.append([sid, k, item.get("id"), json.dumps(item.get("dims") or {}), item.get("label"), p[0], p[1], item.get("unit")])
    _sheet(wb, "ChartData", ["slide", "chart", "series_id", "dims", "label", "period_end", "value", "unit"], rows, [8, 18, 36, 20, 40, 12, 14, 12])

    # KPIs
    rows = [[k, v.get("label"), json.dumps(v.get("dims") or {}), v["latest"]["value"] if v.get("latest") else None, v["latest"]["period"] if v.get("latest") else None,
             v["prior"]["value"] if v.get("prior") else None, v["prior"]["period"] if v.get("prior") else None, v.get("change"), v.get("unit"), v.get("compare"), v.get("period_type"), ",".join(v.get("doc_ids") or [])]
            for k, v in fp["metrics"].items()]
    _sheet(wb, "KPIs", ["ref", "label", "dims", "latest", "latest_period", "prior", "prior_period", "change", "unit", "comparison", "period_type", "doc_ids"], rows, [40, 50, 20, 14, 12, 14, 12, 12, 8, 12, 24, 60])

    # Revisions
    rows = []
    for v in db.all_vintages():
        if v["n_revisions"]:
            for ex in json.loads(v["revision_summary"] or "[]"):
                rows.append([v["vintage_id"], v["dataset_id"], v["created_at"], ex.get("series_id"), json.dumps(ex.get("dims") or {}), ex.get("period_end"), ex.get("old"), ex.get("new")])
    _sheet(wb, "Revisions", ["vintage_id", "dataset_id", "created_at", "series_id", "dims", "period_end", "old_value", "new_value"], rows, [60, 26, 22, 36, 20, 12, 14, 14])

    # Quality
    rows = [[c["check"], c.get("ok"), c.get("n"), c.get("failed"), json.dumps(c.get("examples"), ensure_ascii=False, default=str)[:900], c.get("note")] for c in fp["quality"]["checks"]]
    _sheet(wb, "Quality", ["check", "ok", "n", "failed", "examples", "note"], rows, [46, 6, 6, 8, 90, 50])

    # Availability
    rows = []
    for sid, m in fp["availability"]["slides"].items():
        for i in m["inputs"]:
            rows.append([sid, m["status"], i["input"], i["status"], i.get("latest_period")])
    for u in fp["availability"]["unverified"]:
        rows.append(["—", "unverified", u, "unverified", None])
    _sheet(wb, "Availability", ["slide", "slide_status", "input", "input_status", "latest_period"], rows, [8, 12, 46, 12, 12])

    # Narrative
    rows = [["finding", f.get("id"), f.get("slide_id"), f.get("classification"), f.get("statement"), ",".join(f.get("metric_refs") or []), f.get("period"), f.get("caveat")] for f in nar.get("findings", [])]
    for sid, s in (nar.get("slides") or {}).items():
        rows.append(["slide_title", sid, sid, "title", s.get("title"), "", "", ""])
        for t in s.get("interpretations") or []:
            rows.append(["interpretation", sid, sid, "interpretation", t, "", "", s.get("caveat")])
        rows.append(["so_what", sid, sid, "so_what", s.get("so_what"), "", "", ""])
    val = nar.get("validation") or {}
    rows.append(["validation", "", "", "", f"numbers checked {val.get('numbers_checked')}, problems {len(val.get('problems') or [])}", "", "", json.dumps(val.get("problems"), ensure_ascii=False)[:800]])
    _sheet(wb, "Narrative", ["kind", "id", "slide", "classification", "text", "metric_refs", "period", "caveat"], rows, [14, 8, 8, 18, 100, 40, 12, 60])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
