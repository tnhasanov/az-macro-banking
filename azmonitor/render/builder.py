"""Low-level slide construction on python-pptx with the configured theme.

Design follows the user-supplied reference deck: white slides, purple takeaway title,
"Section | description" subtitle line, lavender panels, purple SO WHAT box, source footer,
logo bottom-right, page number. All charts and tables are native and editable.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any, Sequence

from lxml import etree
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.upper())


class Deck:
    def __init__(self, theme: dict[str, Any], lang: str = "en"):
        self.T = theme
        self.C = theme["colors"]
        self.F = theme["fonts"]
        self.S = theme["sizes"]
        self.L = theme["layout"]
        self.lang = lang
        self.prs = Presentation()
        w, h = theme.get("slide_size", [13.333, 7.5])
        self.prs.slide_width = Inches(w)
        self.prs.slide_height = Inches(h)
        self.W, self.H = w, h
        self._blank = self.prs.slide_layouts[6]
        self.page = 0
        logo = theme.get("logo")
        self.logo_path = str(Path(theme.get("_root", ".")) / logo) if logo else None
        if self.logo_path and not Path(self.logo_path).exists():
            self.logo_path = None

    # ------------------------------------------------------------------ basics
    def new_slide(self):
        self.page += 1
        return self.prs.slides.add_slide(self._blank)

    def add_rect(self, slide, x, y, w, h, fill: str | None = None, line: str | None = None, radius: float | None = 0.06, line_w: float = 1.0):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shp.adjustments[0] = radius
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        shp.text_frame.text = ""
        return shp

    def add_text(self, slide, x, y, w, h, text, *, size: float | None = None, bold: bool = False, color: str | None = None, font: str | None = None,
                 align: str = "l", anchor: str = "t", italic: bool = False, margin: float = 0.0, line_spacing: float | None = None,
                 space_after: float | None = None, bullets: bool = False, autofit: bool = False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if autofit:
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
        tf.vertical_anchor = ANCHOR[anchor]
        paragraphs = text if isinstance(text, list) else [text]
        for i, para in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            if line_spacing:
                p.line_spacing = line_spacing
            if space_after is not None:
                p.space_after = Pt(space_after)
            runs = para if isinstance(para, list) else [{"text": para}]
            if bullets and runs and isinstance(runs[0], dict) and not runs[0].get("no_bullet"):
                self._bullet(p)
            for r in runs:
                if isinstance(r, str):
                    r = {"text": r}
                run = p.add_run()
                run.text = r.get("text", "")
                f = run.font
                f.size = Pt(r.get("size", size or self.S["body"]))
                f.bold = r.get("bold", bold)
                f.italic = r.get("italic", italic)
                f.name = r.get("font") or font or (self.F["title"] if r.get("bold", bold) else self.F["body"])
                f.color.rgb = rgb(r.get("color") or color or self.C["text"])
        return tb

    @staticmethod
    def _bullet(paragraph, char: str = "•", indent_emu: int = 171450):
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("marL", str(indent_emu))
        pPr.set("indent", str(-indent_emu))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", char)

    # ------------------------------------------------------------------ composed elements
    def add_title(self, slide, title: str, section: str, desc: str):
        L = self.L
        self.add_text(slide, L["margin_x"], L["title_top"], self.W - 2 * L["margin_x"] - 0.5, L["title_height"], title, size=self.S["title"], bold=True,
                      color=self.C["primary"], anchor="b")
        runs = [{"text": section, "bold": True, "color": self.C["muted"], "size": self.S["subtitle"]},
                {"text": f"   |   {desc}", "color": self.C["muted"], "size": self.S["subtitle"]}]
        self.add_text(slide, L["margin_x"], L["subtitle_top"], self.W - 2 * L["margin_x"] - 0.5, 0.3, [runs], anchor="m")

    def add_panel(self, slide, x, y, w, h, fill: str | None = None, line: str | None = None):
        return self.add_rect(slide, x, y, w, h, fill or self.C["panel"], line or self.C["panel_border"], radius=0.05)

    def add_so_what(self, slide, x, y, w, h, text: str, heading: str = "SO WHAT"):
        self.add_rect(slide, x, y, w, h, self.C["primary"], None, radius=0.04)
        self.add_text(slide, x + 0.18, y + 0.12, w - 0.36, 0.22, heading, size=8.5, bold=True, color=self.C["white"])
        self.add_text(slide, x + 0.18, y + 0.36, w - 0.36, h - 0.46, text, size=self.S["small"], color=self.C["white"], line_spacing=1.05)

    def add_kpi(self, slide, x, y, w, h, value: str, label: str, sub: str | None = None, value_color: str | None = None, value_size: float | None = None,
                label_size: float | None = None):
        self.add_panel(slide, x, y, w, h)
        vs = value_size or (self.S["kpi_number"] if len(value) <= 8 else self.S["kpi_number"] - 6)
        self.add_text(slide, x + 0.14, y + 0.06, w - 0.28, h * 0.5, value, size=vs, bold=True, color=value_color or self.C["primary"], anchor="m")
        self.add_text(slide, x + 0.14, y + h * 0.52, w - 0.28, h * 0.46, [label] + ([sub] if sub else []),
                      size=label_size or (self.S["kpi_label"] - 1), color=self.C["muted"], line_spacing=1.0)

    def add_footer(self, slide, source: str, page_no: int | None = None):
        L = self.L
        self.add_text(slide, L["margin_x"], L["footer_top"], self.W - 2.0, 0.42, source, size=self.S["source"], color=self.C["muted_light"], anchor="t")
        if self.logo_path:
            lw, lh = self.T.get("logo_size_in", [0.52, 0.34])
            slide.shapes.add_picture(self.logo_path, Inches(self.W - L["margin_x"] - lw), Inches(self.H - 0.6), Inches(lw), Inches(lh))
        if page_no is not None:
            self.add_text(slide, self.W - 1.2, self.H - 0.27, 0.75, 0.2, str(page_no), size=8, color=self.C["muted_light"], align="r")

    def add_notes(self, slide, text: str):
        slide.notes_slide.notes_text_frame.text = text

    # ------------------------------------------------------------------ table
    def add_table(self, slide, x, y, w, h, header: Sequence[str], rows: Sequence[Sequence[str]], col_widths: Sequence[float] | None = None,
                  font_size: float | None = None, align: Sequence[str] | None = None, header_fill: str | None = None, zebra: bool = True,
                  bold_rows: Sequence[int] | None = None, highlight_rows: dict[int, str] | None = None, row_height: float | None = None,
                  first_col_color: str | None = None):
        n_rows, n_cols = len(rows) + 1, len(header)
        gs = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = gs.table
        tbl.first_row = True
        tbl.horz_banding = False
        # remove default style banding by clearing tblPr style id
        tblPr = tbl._tbl.tblPr
        for el in list(tblPr):
            if el.tag == qn("a:tableStyleId"):
                tblPr.remove(el)
        fs = font_size or self.S["table"]
        if col_widths:
            total = sum(col_widths)
            for j, cw in enumerate(col_widths):
                tbl.columns[j].width = Inches(w * cw / total)
        rh = Inches(row_height) if row_height else Inches(h / n_rows)
        for i in range(n_rows):
            tbl.rows[i].height = rh
        align = align or (["l"] + ["r"] * (n_cols - 1))
        for j, text in enumerate(header):
            self._cell(tbl.cell(0, j), str(text), fs, True, self.C["white"], header_fill or self.C["table_header"], align[j] if j < len(align) else "r")
        for i, row in enumerate(rows, start=1):
            fill = self.C["white"] if (not zebra or i % 2 == 1) else self.C["table_stripe"]
            if highlight_rows and (i - 1) in highlight_rows:
                fill = highlight_rows[i - 1]
            bold = bool(bold_rows and (i - 1) in bold_rows)
            for j in range(n_cols):
                text = row[j] if j < len(row) else ""
                color = first_col_color if (j == 0 and first_col_color) else self.C["text"]
                self._cell(tbl.cell(i, j), "" if text is None else str(text), fs, bold, color, fill, align[j] if j < len(align) else "r")
        return gs

    def _cell(self, cell, text: str, size: float, bold: bool, color: str, fill: str, align: str):
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.025)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = ALIGN[align]
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = self.F["title"] if bold else self.F["body"]
        run.font.color.rgb = rgb(color)
        self._cell_borders(cell, self.C["table_border"])

    @staticmethod
    def _cell_borders(cell, color: str, width_pt: float = 0.5):
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            for el in tcPr.findall(qn(tag)):
                tcPr.remove(el)
            ln = etree.SubElement(tcPr, qn(tag))
            ln.set("w", str(int(width_pt * 12700)))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            clr = etree.SubElement(sf, qn("a:srgbClr"))
            clr.set("val", color.upper())
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        # borders must precede fill in tcPr order: move solidFill to the end
        fills = tcPr.findall(qn("a:solidFill"))
        for f in fills:
            tcPr.remove(f)
            tcPr.append(f)

    # ------------------------------------------------------------------ charts
    def _style_axes(self, chart, number_format: str, cat_font: float | None = None, val_font: float | None = None, gridlines: bool = True, skip: int = 1):
        cf = self.S["chart_label"]
        chart.font.size = Pt(cf)
        chart.font.name = self.F["body"]
        chart.font.color.rgb = rgb(self.C["muted"])
        try:
            va = chart.value_axis
            va.has_major_gridlines = gridlines
            if gridlines:
                va.major_gridlines.format.line.color.rgb = rgb(self.C["table_border"])
                va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(val_font or cf)
            va.tick_labels.number_format = number_format
            va.tick_labels.number_format_is_linked = False
            va.tick_labels.font.color.rgb = rgb(self.C["muted"])
        except Exception:
            pass
        try:
            ca = chart.category_axis
            ca.tick_labels.font.size = Pt(cat_font or cf)
            ca.tick_labels.font.color.rgb = rgb(self.C["muted"])
            ca.format.line.color.rgb = rgb(self.C["table_border"])
            ca.has_major_gridlines = False
            ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            if skip > 1:
                el = ca._element
                for old in el.findall(qn("c:tickLblSkip")):
                    el.remove(old)
                s = etree.SubElement(el, qn("c:tickLblSkip"))
                s.set("val", str(skip))
                # tickLblSkip must come after c:lblOffset per schema; append order handled by moving before c:noMultiLvlLbl if present
                nm = el.find(qn("c:noMultiLvlLbl"))
                if nm is not None:
                    el.remove(s)
                    nm.addprevious(s)
        except Exception:
            pass

    def _legend(self, chart, show: bool, position=XL_LEGEND_POSITION.BOTTOM):
        chart.has_legend = show
        if show:
            chart.legend.position = position
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(self.S["chart_label"])
            chart.legend.font.color.rgb = rgb(self.C["muted"])

    def add_line_chart(self, slide, x, y, w, h, categories: Sequence[str], series: Sequence[dict[str, Any]], *, number_format: str = "0.0",
                       legend: bool = True, markers: bool = False, label_last: bool = True, skip: int | None = None, width_pt: float = 2.0):
        cd = CategoryChartData()
        cd.categories = list(categories)
        for s in series:
            cd.add_series(s["name"], [None if v is None else float(v) for v in s["values"]])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        self._style_axes(chart, number_format, skip=skip or max(1, len(categories) // 12))
        self._legend(chart, legend and len(series) > 1)
        plot = chart.plots[0]
        for i, s in enumerate(series):
            ser = plot.series[i]
            ser.smooth = False
            ser.format.line.color.rgb = rgb(s.get("color") or self.C["series_primary"])
            ser.format.line.width = Pt(s.get("width", width_pt))
            if s.get("dash"):
                ser.format.line.dash_style = 4  # MSO_LINE.DASH
            if markers:
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 5
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = rgb(s.get("color") or self.C["series_primary"])
                ser.marker.format.line.fill.background()
            if label_last and s["values"]:
                idx = max(i for i, v in enumerate(s["values"]) if v is not None) if any(v is not None for v in s["values"]) else None
                if idx is not None:
                    pt = ser.points[idx]
                    dl = pt.data_label
                    tf = dl.text_frame
                    tf.text = _fmt_label(s["values"][idx], number_format)
                    run = tf.paragraphs[0].runs[0]
                    run.font.size = Pt(self.S["chart_label"])
                    run.font.bold = True
                    run.font.color.rgb = rgb(s.get("color") or self.C["series_primary"])
                    dl.position = XL_LABEL_POSITION.ABOVE
        return gf

    def add_bar_chart(self, slide, x, y, w, h, categories: Sequence[str], series: Sequence[dict[str, Any]], *, stacked: bool = False, horizontal: bool = False,
                      number_format: str = "0.0", legend: bool = True, data_labels: bool = False, gap_width: int = 60, skip: int | None = None,
                      point_colors: Sequence[str | None] | None = None, label_font: float | None = None, overlap: int | None = None):
        if not categories or not any(v is not None for s in series for v in s["values"]):
            # An empty category chart is not valid Office XML, and python-pptx raises on it; one slide
            # without data used to fail the whole deck. Say what is missing in the chart's place.
            self.add_text(slide, x, y, w, min(h, 0.8), "No figures are available for this period.", size=9,
                          color=self.C.get("muted"))
            return None
        cd = CategoryChartData()
        cd.categories = list(categories)
        for s in series:
            cd.add_series(s["name"], [None if v is None else float(v) for v in s["values"]])
        if horizontal:
            ctype = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
        else:
            ctype = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        self._style_axes(chart, number_format, skip=skip or (max(1, len(categories) // 12) if not horizontal else 1))
        self._legend(chart, legend and len(series) > 1)
        plot = chart.plots[0]
        plot.gap_width = gap_width
        if overlap is not None:
            plot.overlap = overlap
        elif stacked:
            plot.overlap = 100
        if horizontal:
            try:
                chart.category_axis.reverse_order = True
            except Exception:
                pass
        for i, s in enumerate(series):
            ser = plot.series[i]
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(s.get("color") or self.C["series_primary"])
            ser.format.line.fill.background()
            ser.invert_if_negative = False
            if s.get("invisible"):
                ser.format.fill.background()
            if point_colors and i == 0:
                for j, c in enumerate(point_colors):
                    if c:
                        pt = ser.points[j]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = rgb(c)
        if data_labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.font.size = Pt(label_font or self.S["chart_label"])
            dl.font.color.rgb = rgb(self.C["text"])
            dl.position = XL_LABEL_POSITION.INSIDE_END if stacked else XL_LABEL_POSITION.OUTSIDE_END
            for i, s in enumerate(series):
                if s.get("invisible"):
                    plot.series[i].data_labels.show_value = False
                    plot.series[i].data_labels.font.size = Pt(1)
        return gf

    def add_bubble_chart(self, slide, x, y, w, h, points: Sequence[dict[str, Any]], *, x_title: str = "", y_title: str = "", number_format: str = "0.0"):
        cd = BubbleChartData()
        for p in points:
            s = cd.add_series(p["name"])
            s.add_data_point(float(p["x"]), float(p["y"]), float(p.get("size") or 1.0))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        chart.font.size = Pt(self.S["chart_label"])
        chart.font.name = self.F["body"]
        chart.font.color.rgb = rgb(self.C["muted"])
        self._legend(chart, False)
        plot = chart.plots[0]
        plot.bubble_scale = 45
        for i, p in enumerate(points):
            ser = plot.series[i]
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(p.get("color") or self.C["series_primary"])
            ser.format.line.color.rgb = rgb(self.C["white"])
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_series_name = True
        dl.show_value = False
        dl.font.size = Pt(self.S["chart_label"] - 1)
        dl.position = XL_LABEL_POSITION.ABOVE
        for ax, title in ((chart.value_axis, y_title), (chart.category_axis, x_title)):
            ax.has_major_gridlines = True
            ax.major_gridlines.format.line.color.rgb = rgb(self.C["table_border"])
            ax.tick_labels.font.size = Pt(self.S["chart_label"])
            ax.tick_labels.number_format = number_format
            ax.tick_labels.number_format_is_linked = False
            if title:
                ax.has_title = True
                ax.axis_title.text_frame.text = title
                ax.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(self.S["chart_label"])
                ax.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(self.C["muted"])
        return gf

    def add_waterfall(self, slide, x, y, w, h, steps: Sequence[dict[str, Any]], *, number_format: str = "#,##0", pos_color: str | None = None, neg_color: str | None = None,
                      total_color: str | None = None):
        """steps: [{label, value, kind: 'total'|'delta'}] rendered as stacked columns with an invisible base."""
        cats, base, pos, neg, tot = [], [], [], [], []
        running = 0.0
        for s in steps:
            cats.append(s["label"])
            v = float(s["value"] or 0.0)
            if s["kind"] == "total":
                running = v
                base.append(0.0)
                tot.append(v)
                pos.append(None)
                neg.append(None)
            else:
                start = running
                running = start + v
                lo, hi = min(start, running), max(start, running)
                base.append(lo)
                tot.append(None)
                pos.append(hi - lo if v >= 0 else None)
                neg.append(hi - lo if v < 0 else None)
        series = [{"name": "base", "values": base, "invisible": True}, {"name": "Total", "values": tot, "color": total_color or self.C["primary"]},
                  {"name": "Increase", "values": pos, "color": pos_color or self.C["positive"]}, {"name": "Decrease", "values": neg, "color": neg_color or self.C["negative"]}]
        gf = self.add_bar_chart(slide, x, y, w, h, cats, series, stacked=True, legend=False, data_labels=False, gap_width=40, skip=1)
        chart = gf.chart
        plot = chart.plots[0]
        # data labels with the signed step value
        for i, s in enumerate(steps):
            v = float(s["value"] or 0.0)
            ser_idx = 1 if s["kind"] == "total" else (2 if v >= 0 else 3)
            pt = plot.series[ser_idx].points[i]
            dl = pt.data_label
            tf = dl.text_frame
            tf.text = f"{v:+,.0f}" if s["kind"] != "total" else f"{v:,.0f}"
            tf.paragraphs[0].runs[0].font.size = Pt(self.S["chart_label"])
            tf.paragraphs[0].runs[0].font.bold = True
            tf.paragraphs[0].runs[0].font.color.rgb = rgb(self.C["text"])
            dl.position = XL_LABEL_POSITION.INSIDE_END if s["kind"] == "total" else XL_LABEL_POSITION.INSIDE_BASE
        return gf

    def save(self, path: Path | str) -> Path:
        p = Path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(p))
        return p


def _fmt_label(v: float, number_format: str) -> str:
    dec = 0 if "0.0" not in number_format else len(number_format.split(".")[-1].rstrip("%;"))
    return f"{v:,.{dec}f}"


def month_labels(isos: Sequence[str], lang: str = "en") -> list[str]:
    from ..util.periods import AZ_MONTHS_TITLE, EN_MONTHS

    months = EN_MONTHS if lang == "en" else AZ_MONTHS_TITLE
    out = []
    for s in isos:
        d = dt.date.fromisoformat(s)
        out.append(f"{months[d.month - 1][:3]}-{d.year % 100:02d}")
    return out


def align_series(series_list: Sequence[dict[str, Any]]) -> tuple[list[str], list[list[float | None]]]:
    """Align fact-pack chart series (points [[iso, value], ...]) on the union of periods."""
    periods = sorted({p[0] for s in series_list for p in s.get("points", [])})
    values = []
    for s in series_list:
        m = {p[0]: p[1] for p in s.get("points", [])}
        values.append([m.get(p) for p in periods])
    return periods, values
