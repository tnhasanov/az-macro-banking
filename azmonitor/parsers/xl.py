"""Uniform access to .xlsx (openpyxl) and .xls (xlrd) workbooks as lists of rows."""
from __future__ import annotations

import datetime as dt
import warnings
from pathlib import Path
from typing import Any


def load_sheets(path: Path) -> dict[str, list[list[Any]]]:
    """Return {sheet_name: rows}; cells are python values (str/float/int/datetime/None)."""
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        import openpyxl

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        out = {}
        for ws in wb.worksheets:
            out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        return out
    if p.suffix.lower() == ".xls":
        import xlrd

        wb = xlrd.open_workbook(p)
        out = {}
        for ws in wb.sheets():
            rows = []
            for i in range(ws.nrows):
                row = []
                for j in range(ws.ncols):
                    c = ws.cell(i, j)
                    v: Any = c.value
                    if c.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = dt.datetime(*xlrd.xldate_as_tuple(v, wb.datemode))
                        except Exception:
                            pass
                    elif c.ctype == xlrd.XL_CELL_EMPTY:
                        v = None
                    elif c.ctype == xlrd.XL_CELL_ERROR:
                        v = "#REF!"
                    row.append(v)
                rows.append(row)
            out[ws.name] = rows
        return out
    raise ValueError(f"unsupported workbook type: {p}")


def cell_ref(sheet: str, row_idx0: int, col_idx0: int) -> str:
    col = ""
    n = col_idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        col = chr(65 + r) + col
    return f"{sheet}!{col}{row_idx0 + 1}"


def find_sheet(sheets: dict[str, list[list[Any]]], name: str) -> tuple[str, list[list[Any]]]:
    if name in sheets:
        return name, sheets[name]
    norm = name.strip().lower()
    for k in sheets:
        if k.strip().lower() == norm:
            return k, sheets[k]
    raise KeyError(f"sheet {name!r} not found; available: {list(sheets)[:10]}")
